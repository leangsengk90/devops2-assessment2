"""Assemble the dissertation from the UP template.

Keeps the template's cover, section breaks, page numbering, styles and TOC field; replaces all
placeholder content. Run with the scratch venv:  python build.py
"""
import copy
import json
import re
import zipfile
from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

import backmatter
import ch1, ch2, ch3, ch4, ch5, ch6
from docxlib import h2, li, p, table, ref, pagebreak, W

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = ROOT.parent / "resources" / "Dissertation_Complete_Chapters_1-6.docx"
A1 = ROOT.parent / "assessment1" / "KAO_Leangseng_ITRM_Assessment1_v2.docx"
OUT = ROOT / "KAO_Leangseng_ITRM_Assessment2_Dissertation.docx"
FIG = ROOT / "figures"

TITLE = "EVALUATING THE EFFECTIVENESS OF TERRAFORM AND GITHUB ACTIONS FOR AWS INFRASTRUCTURE AUTOMATION"
TITLE_CASE = "Evaluating the Effectiveness of Terraform and GitHub Actions for AWS Infrastructure Automation"
R = json.loads((ROOT / "analysis" / "results.json").read_text())

doc = docx.Document(TEMPLATE)
body = doc.element.body
ch = list(body.iterchildren())
assert len(ch) == 556


def set_text(el, text):
    """Put text in the first run of a paragraph and drop the other runs (keeps formatting)."""
    rs = el.findall(qn("w:r"))
    rs[0].findall(qn("w:t"))[0].text = text
    for t in rs[0].findall(qn("w:t"))[1:]:
        rs[0].remove(t)
    for r in rs[1:]:
        el.remove(r)


# ------------------------------------------------------------------ cover (elements 0-19)
cover = {4: "FACULTY OF SCIENCE AND TECHNOLOGY", 5: "DEPARTMENT OF SCIENCE AND TECHNOLOGY",
         7: TITLE,
         10: "A Dissertation Submitted to the Faculty of Science and Technology, University of Puthisastra, in Partial Fulfilment of the Requirements for the Degree of",
         11: "MASTER OF SCIENCE IN INFORMATION TECHNOLOGY", 13: "KAO LEANGSENG", 14: "Student ID: UPS240590",
         16: "Dr. CHEN SOVANN", 17: "OCTOBER 2026"}
for i, t in cover.items():
    set_text(ch[i], t)
body.remove(ch[8])  # title now fits in one paragraph
body.remove(ch[9])  # one spacer fewer so the cover stays on one page

# swap the placeholder logo for the official UP logo used in Assessment 1
blip = ch[1].find(".//" + qn("a:blip"))
part = doc.part.related_parts[blip.get(qn("r:embed"))]
with zipfile.ZipFile(A1) as z:
    part._blob = z.read("word/media/image3.png")


# ------------------------------------------------------------------ front matter
def X(xml):
    return parse_xml(xml)


def fm_h1(text):  # front-matter heading: page break before, as in the template's PB + Heading1 pairs
    return (f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/><w:pageBreakBefore/><w:spacing w:after="240" w:before="240"/>'
            f'<w:jc w:val="left"/></w:pPr><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>')


pooled = R["gh_phases"]["Full create"]
sim = R["sim_manual"]
h1 = R["h1"]
front_a = [
    p(f"I, KAO LEANGSENG, Student ID UPS240590, hereby declare that this dissertation entitled “{TITLE_CASE}” is my own original work, produced under the supervision of Dr. CHEN SOVANN. It has not been submitted, in whole or in part, for any other degree or qualification at the University of Puthisastra or at any other institution."),
    p("All sources of information, ideas, data and prior scholarship used in this dissertation have been acknowledged in the text and listed in full in the reference list, in accordance with the seventh edition of the Publication Manual of the American Psychological Association. I confirm that I have complied with the University of Puthisastra policy on academic integrity, and I accept that any breach of that policy, including plagiarism, falsification of data or unacknowledged use of generative artificial intelligence tools, may result in disciplinary action. The use of generative artificial intelligence in preparing this dissertation is acknowledged on the Acknowledgements page."),
    p(""), p(""),
    p("Signature: ______________________                    Date: ______________________"),
    p("Name: KAO LEANGSENG"),
    fm_h1("SUPERVISOR’S APPROVAL"),
    p(f"This is to certify that the dissertation entitled “{TITLE_CASE}”, submitted by KAO LEANGSENG (ID: UPS240590) to the Faculty of Science and Technology, University of Puthisastra, has been prepared under my supervision. I confirm that the work meets the academic standards required for submission in partial fulfilment of the requirements for the degree of Master of Science in Information Technology, and I recommend it for examination."),
    p(""), p(""),
    p("Signature: ______________________                    Date: ______________________"),
    p("Supervisor: Dr. CHEN SOVANN"),
    p("Faculty of Science and Technology, University of Puthisastra"),
    fm_h1("ACKNOWLEDGEMENTS"),
    p("I want to thank my supervisor, Dr. CHEN SOVANN, whose guidance and honest feedback shaped this research at every stage. I am also grateful to the teaching and administrative staff of the Faculty of Science and Technology at the University of Puthisastra for the training that made this study possible."),
    p("I thank the management of the case organisation for permission to use its cloud architecture as the basis of this study and for approving a separate sandbox account so that the work could be done without any risk to live systems, and my colleagues in the infrastructure team, whose day-to-day questions about provisioning first made me notice the problem this dissertation addresses."),
    p("Generative artificial intelligence (Claude, by Anthropic) was used as an assistant in this project: to help write and debug parts of the GitHub Actions workflow and measurement scripts, to run the statistical analysis scripts, to help verify references, and to help draft and edit the text of this dissertation. All data reported were produced by the infrastructure and scripts described in Chapter 3, and I have reviewed and take responsibility for all content."),
    p("Finally, I thank my family and friends for their patience and encouragement. Any errors that remain are my own."),
    fm_h1("ABSTRACT"),
    p("**Background.** Infrastructure as Code (IaC) is standard practice in cloud engineering, yet many organisations automate application delivery long before the infrastructure it runs on. Existing research compares IaC tools with each other or studies IaC code quality in isolation, but does not measure what changes when Terraform is introduced into a CI/CD pipeline in place of manual provisioning, and offers no evidence from Cambodia."),
    p("**Aim.** The study designed, implemented and evaluated a modular Terraform codebase for a five-service containerised AWS architecture, integrated with GitHub Actions, and measured provisioning duration, configuration drift, reproducibility and version-control coverage, while explaining the barriers encountered."),
    p(f"**Methods.** A pragmatist, explanatory sequential mixed methods design was used. The quantitative strand measured {R['n_tf']} complete destroy-and-rebuild cycles of a 152-resource environment executed by GitHub Actions over two days, {R['n_starts']} service start-ups and a 24-item component audit. Manual creation of the same architecture in the AWS console was represented by a seeded Monte Carlo simulation of 10 builds and 10 rebuilds, with every parameter stated and tested for sensitivity. Data were analysed with descriptive statistics, stability checks, Welch's _t_, Mann\u2013Whitney _U_, Hedges' _g_ and Fisher's exact test. The qualitative strand applied reflexive thematic analysis to {R['record']['commits']} commits, {R['record']['run_reports']} run reports and the pipeline code."),
    p(f"**Results.** Terraform executed by GitHub Actions built the environment from empty in a mean of {pooled['mean']/60:.1f} minutes (_M_ = {pooled['mean']:.1f} s, _SD_ = {pooled['sd']:.1f}) and destroyed it in {R['gh_phases']['Full destroy']['mean']/60:.1f} minutes; all {R['n_tf']} rebuilds and all {R['n_starts']} service start-ups succeeded, with no trend over run order, day or hour of day. The simulated manual build took {sim['mean']/60:.1f} minutes, {h1['ratio']:.1f} times longer (Welch's _t_({h1['welch_df']:.1f}) = {h1['welch_t']:.2f}, _p_ < .001), and only {R['man_rebuild_success']['k']} of 10 simulated manual rebuilds were identical to the baseline against {R['n_tf']} of {R['n_tf']} for Terraform (Fisher's exact _p_ {'< .001' if R['h3_fisher_p'] < .001 else '= ' + format(R['h3_fisher_p'], '.3f')}); these comparisons are conditional on the simulation. Version-control coverage rose from 0% to 79.2%, but no change was peer-reviewed. Drift could not be measured because only one environment existed, so H2 and H5 were not tested. Four themes explained the results: the architecture converged through repeated rebuilding, state management was the hardest problem, destruction rather than creation was the bottleneck, and version control arrived without review."),
    p("**Conclusion.** Terraform with GitHub Actions made the architecture fast and reliable to rebuild compared with manual console work, but review did not follow automatically. For regulated Cambodian institutions, the code history must be paired with an enforced review gate to provide controlled change. The study contributes measured benchmarks, a reusable codebase and procedure, and the first evidence of this kind from Cambodia."),
    p(""),
    p("**Keywords:** Infrastructure as Code; Terraform; GitHub Actions; continuous integration and continuous delivery; Amazon Web Services; reproducibility; mixed methods; Cambodia"),
    pagebreak(),
]

# replace elements 21..49 (declaration body .. abstract) with front_a; keep 50 (TOC heading) and 51 (TOC field)
for el in ch[21:50]:
    body.remove(el)
anchor = ch[50]
for xml in front_a:
    anchor.addprevious(X(xml))

# ------------------------------------------------------------------ chapters + back matter
blocks = ch1.BLOCKS + ch2.BLOCKS + ch3.BLOCKS + ch4.BLOCKS + ch5.BLOCKS + ch6.BLOCKS
blocks += [X_ for X_ in []]
refs = ['<REFHEAD/>'] + [ref(r) for r in backmatter.REFS]
appx = backmatter.APPX + backmatter.appx_stats(R["gh_cycles"], R["sim_manual_runs"])

# list of tables / figures from captions
tables, figures = [], []
for b in blocks + appx:
    if isinstance(b, str):
        m = re.search(r'<w:t xml:space="preserve">(Table [0-9]+\.[0-9]+) </w:t></w:r>(.*?)</w:p>', b)
        if m:
            title = re.sub(r"<[^>]+>", "", m.group(2))
            tables.append(f"{m.group(1)}  {title}")
    else:
        pass
for b in blocks:
    if isinstance(b, str) and "Figure " in b and 'w:sz w:val="20"' in b and "w:after=\"220\"" in b:
        t = re.sub(r"<[^>]+>", "", b)
        m = re.match(r"(Figure [0-9]+\.[0-9]+)\. ([^.]+)", t)
        figures.append(f"{m.group(1)}  {m.group(2)}")

lists = [fm_h1("LIST OF TABLES")] + [p(t, after=100) for t in tables]
lists += [fm_h1("LIST OF FIGURES")] + [p(f, after=100) for f in figures]
lists += [fm_h1("LIST OF ABBREVIATIONS"), table(["Abbreviation", "Full Term"], [
    ["ALB", "Application Load Balancer"], ["API", "Application Programming Interface"],
    ["APA", "American Psychological Association"], ["AWS", "Amazon Web Services"], ["AZ", "Availability Zone"],
    ["CI", "Confidence Interval (statistics) / Continuous Integration (engineering)"],
    ["CI/CD", "Continuous Integration and Continuous Delivery"], ["CLI", "Command Line Interface"],
    ["CV", "Coefficient of Variation"], ["DNS", "Domain Name System"], ["DOI", "Digital Object Identifier"],
    ["ECR", "Elastic Container Registry"], ["ECS", "Elastic Container Service"], ["GHA", "GitHub Actions"],
    ["HTTP", "Hypertext Transfer Protocol"], ["IaC", "Infrastructure as Code"],
    ["IAM", "Identity and Access Management"], ["NAT", "Network Address Translation"],
    ["OIDC", "OpenID Connect"], ["RQ", "Research Question"], ["S3", "Simple Storage Service"],
    ["SCP", "Service Control Policy"], ["SD", "Standard Deviation"], ["TLS", "Transport Layer Security"],
    ["UP", "University of Puthisastra"], ["VPC", "Virtual Private Cloud"], ["WAF", "Web Application Firewall"],
    ["YAML", "YAML Ain't Markup Language"]], [2200, 6440])]

for el in ch[52:82]:
    body.remove(el)
sect_body = ch[82]
for xml in lists:
    sect_body.addprevious(X(xml))

final_sect = ch[555]
for el in ch[83:555]:
    body.remove(el)


body_sect = body.find(qn("w:sectPr"))


def sect_par(landscape, keep_start):
    s = copy.deepcopy(body_sect)
    pg = s.find(qn("w:pgNumType"))
    if pg is not None and not keep_start and pg.get(qn("w:start")) is not None:
        del pg.attrib[qn("w:start")]
    if landscape:
        sz = s.find(qn("w:pgSz"))
        w, h = sz.get(qn("w:w")), sz.get(qn("w:h"))
        sz.set(qn("w:w"), h); sz.set(qn("w:h"), w); sz.set(qn("w:orient"), "landscape")
        mar = s.find(qn("w:pgMar"))
        mar.set(qn("w:left"), "1440"); mar.set(qn("w:right"), "1440")
    par = X(f'<w:p {W}><w:pPr></w:pPr></w:p>')
    par.find(qn("w:pPr")).append(s)
    return par


def add(b):
    if isinstance(b, tuple) and b[0] == "landscape_start":
        final_sect.addprevious(sect_par(False, True))   # ends the portrait section that began at Chapter 1
    elif isinstance(b, tuple) and b[0] == "landscape_end":
        final_sect.addprevious(sect_par(True, False))   # ends the landscape section
        pg = body_sect.find(qn("w:pgNumType"))
        if pg is not None and pg.get(qn("w:start")) is not None:
            del pg.attrib[qn("w:start")]                # remaining pages continue the numbering
    elif isinstance(b, tuple) and b[0] == "fig":
        par = X(f'<w:p {W}><w:pPr><w:keepNext/><w:spacing w:after="80" w:before="200" w:line="240"/><w:jc w:val="center"/></w:pPr></w:p>')
        final_sect.addprevious(par)
        docx.text.paragraph.Paragraph(par, doc._body).add_run().add_picture(str(FIG / b[1]), width=Inches(b[2]))
    elif b == "<REFHEAD/>":
        final_sect.addprevious(X(fm_h1("REFERENCES")))
        final_sect.addprevious(X(p("All sources are formatted in accordance with the seventh edition of the Publication Manual of the American Psychological Association (American Psychological Association, 2020) and are listed alphabetically by author surname with a hanging indent of 0.5 inch.", after=240)))
    else:
        final_sect.addprevious(X(b))


for b in blocks + refs + appx:
    add(b)

doc.core_properties.title = TITLE_CASE
doc.core_properties.author = "KAO LEANGSENG"
doc.save(OUT)
print("saved", OUT.name, "tables", len(tables), "figures", len(figures))
